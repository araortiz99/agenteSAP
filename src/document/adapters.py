"""Bounded read-only adapters for PDF, DOCX and XLSX sources.

Adapters normalize supported binary office/document formats into DocumentSource.
They never write to disk, execute macros, or access SAP/MCP.
"""

from __future__ import annotations

from pathlib import Path
from typing import Callable

from src.document.source import DocumentSource

SUPPORTED_BINARY_TYPES = frozenset({"pdf", "docx", "xlsx"})

_MEDIA_TYPES = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


def _validate_path(path: str | Path, max_bytes: int, expected: str) -> Path:
    if max_bytes < 1:
        raise ValueError("max_bytes must be greater than zero")
    file_path = Path(path)
    if not file_path.is_file():
        raise FileNotFoundError(str(file_path))
    suffix = file_path.suffix.lower().lstrip(".")
    if suffix != expected:
        raise ValueError(f"expected .{expected} file, got .{suffix or 'unknown'}")
    if file_path.stat().st_size > max_bytes:
        raise ValueError(f"file exceeds max_bytes={max_bytes}")
    return file_path


def _source(path: Path, content: str, *, media_type: str) -> DocumentSource:
    return DocumentSource(
        source_id=path.name,
        filename=path.name,
        content=content,
        media_type=media_type,
        metadata=(("source", "LOCAL_FILE"), ("adapter", path.suffix.lower().lstrip("."))),
    )


def load_pdf_file(
    path: str | Path,
    *,
    max_bytes: int = 10_000_000,
    max_pages: int = 64,
    max_chars: int = 200_000,
) -> DocumentSource:
    """Extract bounded text from a PDF using pypdf."""
    file_path = _validate_path(path, max_bytes, "pdf")
    if max_pages < 1 or max_chars < 1:
        raise ValueError("max_pages and max_chars must be greater than zero")
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError("PDF adapter requires the pypdf dependency") from exc

    reader = PdfReader(str(file_path))
    parts: list[str] = []
    total = 0
    for page in reader.pages[:max_pages]:
        text = page.extract_text() or ""
        remaining = max_chars - total
        if remaining <= 0:
            break
        text = text[:remaining]
        if text:
            parts.append(text)
            total += len(text)
    return _source(file_path, "\n\n".join(parts), media_type=_MEDIA_TYPES["pdf"])


def load_docx_file(
    path: str | Path,
    *,
    max_bytes: int = 10_000_000,
    max_paragraphs: int = 5_000,
    max_chars: int = 200_000,
) -> DocumentSource:
    """Extract bounded paragraph/table text from a DOCX without executing macros."""
    file_path = _validate_path(path, max_bytes, "docx")
    if max_paragraphs < 1 or max_chars < 1:
        raise ValueError("max_paragraphs and max_chars must be greater than zero")
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("DOCX adapter requires the python-docx dependency") from exc

    document = Document(str(file_path))
    parts: list[str] = []
    total_items = 0
    total_chars = 0

    def add_text(value: str) -> bool:
        nonlocal total_items, total_chars
        if total_items >= max_paragraphs or total_chars >= max_chars:
            return False
        value = value.strip()
        if not value:
            return True
        remaining = max_chars - total_chars
        parts.append(value[:remaining])
        total_chars += min(len(value), remaining)
        total_items += 1
        return total_chars < max_chars and total_items < max_paragraphs

    for paragraph in document.paragraphs:
        if not add_text(paragraph.text):
            break
    if total_items < max_paragraphs and total_chars < max_chars:
        for table in document.tables:
            for row in table.rows:
                if not add_text(" | ".join(cell.text.strip() for cell in row.cells)):
                    break
            if total_items >= max_paragraphs or total_chars >= max_chars:
                break

    return _source(file_path, "\n\n".join(parts), media_type=_MEDIA_TYPES["docx"])


def load_xlsx_file(
    path: str | Path,
    *,
    max_bytes: int = 10_000_000,
    max_rows: int = 5_000,
    max_cells: int = 50_000,
    max_chars: int = 200_000,
) -> DocumentSource:
    """Extract bounded worksheet cell values from XLSX in read-only mode."""
    file_path = _validate_path(path, max_bytes, "xlsx")
    if min(max_rows, max_cells, max_chars) < 1:
        raise ValueError("max_rows, max_cells and max_chars must be greater than zero")
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("XLSX adapter requires the openpyxl dependency") from exc

    workbook = load_workbook(
        filename=str(file_path),
        read_only=True,
        data_only=True,
    )
    parts: list[str] = []
    cells = 0
    total_chars = 0
    try:
        for worksheet in workbook.worksheets:
            rows = 0
            for row in worksheet.iter_rows(values_only=True):
                if rows >= max_rows or cells >= max_cells or total_chars >= max_chars:
                    break
                values = ["" if value is None else str(value) for value in row]
                if not any(value.strip() for value in values):
                    rows += 1
                    continue
                line = " | ".join(values)
                remaining = max_chars - total_chars
                line = line[:remaining]
                parts.append(f"[{worksheet.title}] {line}")
                total_chars += len(line) + 1
                cells += len(values)
                rows += 1
            if rows >= max_rows or cells >= max_cells or total_chars >= max_chars:
                break
    finally:
        workbook.close()

    return _source(file_path, "\n".join(parts), media_type=_MEDIA_TYPES["xlsx"])


def load_supported_binary_file(path: str | Path, **kwargs) -> DocumentSource:
    """Dispatch one supported binary format to its bounded adapter."""
    suffix = Path(path).suffix.lower().lstrip(".")
    loaders: dict[str, Callable[..., DocumentSource]] = {
        "pdf": load_pdf_file,
        "docx": load_docx_file,
        "xlsx": load_xlsx_file,
    }
    try:
        loader = loaders[suffix]
    except KeyError as exc:
        raise ValueError(
            f"unsupported binary document type: .{suffix or 'unknown'}; "
            f"supported types: {', '.join(sorted(SUPPORTED_BINARY_TYPES))}"
        ) from exc
    return loader(path, **kwargs)
